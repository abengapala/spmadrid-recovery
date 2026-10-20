"""Recovery Status Report pipeline: report + DRR + whole field -> updated report."""
from __future__ import annotations
import copy, re
import pandas as pd
import openpyxl

from core.file_io import decrypt_workbook, load_sheet_as_df, load_field_file, DEFAULT_PASSWORD
from core.load2 import load_drr_clean
from core.rebuild2 import rebuild_status_history, rebuild_fv
from core.output import encrypt_workbook
from core.remarks import (select_daily_action_code, is_held, apply_address_status_placeholder,
                          update_status_columns, FV_PLACEHOLDER, parse_remarks_string)

ROUTINE = ["DATE", "ACTION CODE", "STATUS REMARKS", "FV REMARKS", "CLIENT STATUS",
           "ADDRESS STATUS", "UNIT STATUS", "RFD"]
STATUS_COLS = ["CLIENT STATUS", "ADDRESS STATUS", "UNIT STATUS", "RFD"]
KW = re.compile(r"PTP|NEGO|KEPT", re.I)
FIXED = dict(fin="UHNDIMAYACYAC", mode="REV", person="OTH", place="OTH", by="MARJELVIN DIAZ")


def npn(v):
    try:
        f = float(str(v)); return str(int(f))
    except Exception:
        return str(v).strip()


def run(report_bytes, drr_bytes, field_bytes, password=DEFAULT_PASSWORD,
        report_date=None, conflict_mode="newer", reference_bytes=None):
    """conflict_mode: 'newer' (field only overrides when newer than your latest FV entry),
    'field' (field always wins), 'mine' (never overwrite a filled cell)."""
    warn, notes = [], {}
    daily = load_sheet_as_df(decrypt_workbook(report_bytes, password), "DAILY")
    daily["PN"] = daily["PN"].apply(npn)
    ptp_inv = load_sheet_as_df(decrypt_workbook(report_bytes, password), "PTP INVENTORY")
    ptp_inv["PN"] = ptp_inv["PN"].apply(npn)
    from core.rawdrr import load_reference
    ref = load_reference(reference_bytes) if reference_bytes else None
    drr = load_drr_clean(drr_bytes, warn, ref)
    field, fw = load_field_file(field_bytes); warn += fw
    if drr.empty or drr["date_parsed"].isna().all():
        raise ValueError("DRR has no usable dated rows")
    cutoff = drr["date_parsed"].min().normalize()
    RD = pd.Timestamp(report_date) if report_date else drr["date_parsed"].max().normalize()
    notes.update(report_date=RD, cutoff=cutoff, drr_dates=sorted({d.date() for d in drr["date_parsed"]}))

    out = daily.copy()
    book = set(out["PN"]); invpn = set(ptp_inv["PN"])
    stats, zero, status_changes, skipped, fv_changed, from_ph = {}, [], [], [], [], []
    for idx, row in out.iterrows():
        pn = row["PN"]
        if is_held(row.get("PULLED_OUT TAG")):
            continue
        a = drr[drr["account_no"] == pn]
        out.at[idx, "STATUS REMARKS"] = rebuild_status_history(str(row["STATUS REMARKS"] or ""), a, cutoff, stats)
        win = a[a["date_parsed"] >= cutoff]
        if win.empty:
            zero.append(pn)
        else:
            ac = select_daily_action_code(win)
            if ac:
                out.at[idx, "ACTION CODE"] = ac
        af = field[field["pn"] == pn]
        cur_fv = str(row["FV REMARKS"] or "").strip()
        if not af.empty:
            new_fv, ch = rebuild_fv(cur_fv, af, stats)
            if ch:
                out.at[idx, "FV REMARKS"] = new_fv; fv_changed.append(pn)
                if cur_fv in ("", FV_PLACEHOLDER): from_ph.append(pn)
            latest = af.sort_values("date_parsed", ascending=False).iloc[0]
            pd_ = [e.date for e in parse_remarks_string(cur_fv) if e.date is not None]
            prev = max(pd_) if pd_ else None
            newer = prev is None or (pd.notna(latest["date_parsed"]) and latest["date_parsed"] > prev)
            upd = update_status_columns(dict(out.loc[idx]), latest)
            for c in STATUS_COLS:
                cv = out.at[idx, c]
                blank = cv is None or str(cv).strip() in ("", "nan", "None")
                if str(upd.get(c) or "") == str(cv or ""):
                    continue
                allow = blank or conflict_mode == "field" or (conflict_mode == "newer" and newer)
                (status_changes if allow else skipped).append((pn, c, cv, upd[c]))
                if allow: out.at[idx, c] = upd[c]
        fixed = apply_address_status_placeholder(dict(out.loc[idx]))
        out.at[idx, "ADDRESS STATUS"] = fixed["ADDRESS STATUS"]
        out.at[idx, "DATE"] = RD

    # ---- Trails: effort on report date + PTP/NEGO/KEPT status
    rows = []
    for _, r in out.iterrows():
        pn = r["PN"]
        if is_held(r.get("PULLED_OUT TAG")): continue
        a = drr[drr["account_no"] == pn]
        today = a[a["date_parsed"] == RD].sort_values("time_str")
        hit = bool(KW.search(str(r["ACTION CODE"] or ""))) or pn in invpn
        if today.empty and not hit: continue
        amt = float("nan"); reason = "effort"
        if not today.empty:
            pk = today[today["action_code"].str.contains(KW) | today["remark"].str.contains(KW)]
            e = (pk if not pk.empty else today).iloc[-1]
            d, txt, code, card, amt = RD, e["remark"], e["action_code"], e["card_no"], e["ptp_amount"]
            if hit: reason = "effort + PTP/NEGO status"
        else:
            reason = "PTP/NEGO status only (older entry)"
            w = a[a["action_code"].str.contains(KW)].sort_values(["date_parsed", "time_str"])
            if not w.empty:
                e = w.iloc[-1]; d, txt, code, card, amt = e["date_parsed"], e["remark"], e["action_code"], e["card_no"], e["ptp_amount"]
            else:
                ents = [x for x in parse_remarks_string(str(r["STATUS REMARKS"])) if x.date is not None]
                pk = [x for x in ents if KW.search(x.text)] or ents
                if not pk: continue
                d, txt, code, card = pk[0].date, pk[0].text, str(r["ACTION CODE"] or ""), ""
        if not card:
            c = a["card_no"].replace("", pd.NA).dropna(); card = c.iloc[0] if len(c) else None
        rows.append(dict(pn=pn, name=r["NAME"], d=d, txt=txt, code=code, card=card, amt=amt, reason=reason))
    trails = pd.DataFrame(rows).sort_values("pn").reset_index(drop=True)

    # ---- PTP / NEGO candidates
    win_all = drr[drr["date_parsed"] >= cutoff]
    p = win_all[win_all["action_code"].str.contains(KW) | win_all["remark"].str.contains(r"\bKEPT\b|\bPTP\b", case=False)]
    cand = p.groupby("account_no").agg(codes=("action_code", lambda s: ", ".join(sorted(set(s)))),
                                       last=("date_parsed", "max")).reset_index()
    cand["on DAILY"] = cand["account_no"].isin(book); cand["in PTP INVENTORY"] = cand["account_no"].isin(invpn)

    # ---- PTP INVENTORY: any DAILY account now PTP / NEGO / KEPT that is not in the inventory yet
    PROMISE = re.compile(r"PTP|NEGO|KEPT", re.I)
    new_ptp = out[(out["ACTION CODE"].astype(str).apply(lambda x: bool(PROMISE.search(x))))
                  & ~out["PN"].isin(invpn)
                  & ~out["PULLED_OUT TAG"].apply(is_held)].copy()

    # ---- write workbook in place
    wb = openpyxl.load_workbook(decrypt_workbook(report_bytes, password), data_only=False)
    ws = wb["DAILY"]; H = {c.value: i + 1 for i, c in enumerate(ws[1])}
    o = out.set_index("PN")
    for r in range(2, ws.max_row + 1):
        pn = npn(ws.cell(r, H["PN"]).value)
        if pn not in o.index: continue
        for c in ROUTINE:
            if c not in H: continue
            v = o.at[pn, c]
            if isinstance(v, float) and pd.isna(v): v = None
            if isinstance(v, pd.Timestamp): v = v.to_pydatetime()
            if ws.cell(r, H[c]).value != v: ws.cell(r, H[c]).value = v
    tw = wb["Trails Upload"]; TH = {c.value: i + 1 for i, c in enumerate(tw[1])}
    style = [copy.copy(tw.cell(2, c)._style) for c in range(1, len(TH) + 1)]
    old = {npn(tw.cell(r, 2).value): (tw.cell(r, TH["CURRENCY"]).value, tw.cell(r, TH["ACTION AMOUNT"]).value)
           for r in range(2, tw.max_row + 1)}
    for r in range(tw.max_row, 1, -1):
        for c in range(1, len(TH) + 1): tw.cell(r, c).value = None
    fmt = lambda x: x.strftime("%m.%d.%Y")
    for i, x in trails.iterrows():
        r = i + 2
        for c in range(1, len(TH) + 1): tw.cell(r, c)._style = copy.copy(style[c - 1])
        cur = amt = None
        if x.pn in old and (old[x.pn][0] or old[x.pn][1]): cur, amt = old[x.pn]
        elif pd.notna(x.amt): cur, amt = "PHP", f"{x.amt:,.2f}"
        vals = {"FINANCIER ID": FIXED["fin"], "APPLICATION ID": x.pn, "CUSTOMER ID": x.card if pd.notna(x.card) else None,
                "ACTION DATE": fmt(x.d), "ACTION CODE": x.code, "CONTACT MODE": FIXED["mode"],
                "PERSON CONTACTED": FIXED["person"], "PLACE CONTACTED": FIXED["place"], "CURRENCY": cur,
                "ACTION AMOUNT": amt, "NEXT ACTION DATE": fmt(RD + pd.Timedelta(days=1)),
                "CONTACTED BY": FIXED["by"], "REMARKS": f"{fmt(x.d)} {x.txt}"}
        for k, v in vals.items():
            if k in TH: tw.cell(r, TH[k]).value = v
    pw_ = wb["PTP INVENTORY"]; PH = {c.value: i + 1 for i, c in enumerate(pw_[1])}
    last = max([r for r in range(2, pw_.max_row + 1) if pw_.cell(r, PH["PN"]).value not in (None, "")] or [1])
    for _, x in new_ptp.iterrows():
        last += 1
        for c in range(1, len(PH) + 1): pw_.cell(last, c)._style = copy.copy(pw_.cell(last - 1, c)._style)
        for k, ci in PH.items():
            if k in x.index:
                v = x[k]
                if isinstance(v, float) and pd.isna(v): v = None
                if isinstance(v, pd.Timestamp): v = v.to_pydatetime()
                pw_.cell(last, ci).value = v
    data = encrypt_workbook(wb, password)

    # ---- checks
    ends = pd.to_datetime(o.loc[zero, "ENDO DATE"], errors="coerce") if zero else pd.Series(dtype="datetime64[ns]")
    new_endo = [pn for pn in zero if ends[pn] >= RD - pd.Timedelta(days=14)]
    no_eff = o.loc[[pn for pn in zero if pn not in new_endo], ["NAME", "ENDO DATE", "ACTION CODE"]].reset_index()
    nofield = o.loc[[pn for pn in o.index if field[field["pn"] == pn].empty], ["NAME", "ADDRESS STATUS"]].reset_index()
    chg = pd.DataFrame(status_changes, columns=["PN", "COLUMN", "OLD", "NEW"])
    sk = pd.DataFrame(skipped, columns=["PN", "COLUMN", "YOUR VALUE", "FIELD VALUE"])
    sk = sk[sk["FIELD VALUE"].astype(str).str.strip().isin(["", "0", "nan", "None"]) == False]
    return dict(file=data, notes=notes, warnings=warn, stats=stats, accounts=len(out), trails=trails,
                no_effort=no_eff, new_endo=new_endo, no_field=nofield, changed=chg, conflicts=sk,
                candidates=cand, new_ptp=new_ptp[[c for c in ['PN','NAME','ACTION CODE','STATUS REMARKS'] if c in new_ptp.columns]], fv_changed=len(fv_changed), from_placeholder=len(from_ph))


def precheck(report_bytes, drr_bytes, field_bytes, password=DEFAULT_PASSWORD, reference_bytes=None):
    """Quick look at the three files BEFORE running: nothing is changed."""
    from core.rawdrr import load_reference
    ref = load_reference(reference_bytes) if reference_bytes else None
    warn = []
    daily = load_sheet_as_df(decrypt_workbook(report_bytes, password), "DAILY")
    drr = load_drr_clean(drr_bytes, warn, ref)
    field, fw = load_field_file(field_bytes)
    last_asof = pd.to_datetime(daily["DATE"], errors="coerce").max()
    dmin, dmax = drr["date_parsed"].min(), drr["date_parsed"].max()
    expected = pd.bdate_range(last_asof + pd.Timedelta(days=1), dmax) if pd.notna(last_asof) else []
    have = set(drr["date_parsed"].dt.normalize())
    missing = [d.date() for d in expected if d not in have]
    return dict(accounts=len(daily), report_asof=last_asof, drr_from=dmin, drr_to=dmax, drr_rows=len(drr),
                missing_dates=missing, field_rows=len(field), field_latest=field["date_parsed"].max(),
                field_unmatched=sum("non-account-number" in w for w in fw), warnings=warn + fw)
